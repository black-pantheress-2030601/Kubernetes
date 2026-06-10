#!/usr/bin/env python3
"""
Generate PowerPoint presentation for EKS Multi-Tenant Cluster
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Define colors
    AWS_ORANGE = RGBColor(255, 153, 0)
    DARK_GRAY = RGBColor(35, 47, 62)
    LIGHT_GRAY = RGBColor(149, 165, 166)
    GREEN = RGBColor(39, 174, 96)
    RED = RGBColor(231, 76, 60)
    BLUE = RGBColor(52, 152, 219)

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    left = Inches(1)
    top = Inches(2.5)
    width = Inches(8)
    height = Inches(1)

    title_box = slide.shapes.add_textbox(left, top, width, height)
    title_frame = title_box.text_frame
    title_frame.text = "Multi-Tenant EKS Cluster"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = DARK_GRAY
    title_para.alignment = PP_ALIGN.CENTER

    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(3.5), Inches(8), Inches(0.5))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Current State & Future Roadmap"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = LIGHT_GRAY
    subtitle_para.alignment = PP_ALIGN.CENTER

    date_box = slide.shapes.add_textbox(Inches(1), Inches(6.5), Inches(8), Inches(0.5))
    date_frame = date_box.text_frame
    date_frame.text = "June 2026 • AWS Access Account"
    date_para = date_frame.paragraphs[0]
    date_para.font.size = Pt(14)
    date_para.font.color.rgb = LIGHT_GRAY
    date_para.alignment = PP_ALIGN.CENTER

    # Slide 2: Executive Summary
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Executive Summary"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "What We've Built:"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = DARK_GRAY

    items = [
        "✅ Production-ready multi-tenant EKS cluster",
        "✅ Namespace isolation with RBAC",
        "✅ Separate node pools (system vs. tenant workloads)",
        "✅ Persistent storage with EBS volumes (tested & verified)",
        "✅ Automated node updates (Bottlerocket + BRUPOP)",
        "✅ Pod Identity for secure AWS service access"
    ]

    for item in items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.level = 1
        if "✅" in item:
            p.font.color.rgb = GREEN

    p = tf.add_paragraph()
    p.text = "\nValue Delivered:"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = DARK_GRAY

    value_items = [
        "Cost efficiency through resource sharing",
        "Strong tenant isolation (namespace + compute + identity)",
        "Self-service capabilities for teams"
    ]

    for item in value_items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(16)
        p.level = 1

    # Slide 3: Architecture Overview
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Architecture Overview"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    arch_text = """┌─────────────────────────────────────────┐
│         EKS Control Plane               │
│         (Managed by AWS)                │
└──────────────┬──────────────────────────┘
               │
    ┌──────────┴──────────┐
    │                     │
┌───▼────────┐    ┌──────▼──────────┐
│System Nodes│    │ Tenant Nodes    │
│ (no taints)│    │  (tainted)      │
│            │    │                 │
│ • BRUPOP   │    │ • customer1     │
│ • CoreDNS  │    │ • customer2     │
│ • EBS CSI  │    │ • customer3     │
│ • Add-ons  │    │                 │
└────────────┘    └─────────────────┘"""

    p = tf.paragraphs[0]
    p.text = arch_text
    p.font.size = Pt(13)
    p.font.name = 'Courier New'
    p.font.color.rgb = DARK_GRAY

    p = tf.add_paragraph()
    p.text = "\nKey Components:"
    p.font.size = Pt(18)
    p.font.bold = True

    components = [
        "System Nodes: Platform services (BRUPOP, controllers, add-ons)",
        "Tenant Nodes: Workload-specific, tainted to prevent co-location",
        "Bottlerocket OS: Minimal, immutable, auto-updating"
    ]

    for item in components:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.level = 1

    # Slide 4: Multi-Tenancy Design
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Multi-Tenancy Design"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "Isolation Layers:"
    p.font.size = Pt(20)
    p.font.bold = True

    layers = [
        ("Namespace", "tenant-customer1/2/3", "Logical boundary"),
        ("Compute", "Node taints + tolerations", "Physical isolation"),
        ("Network", "NetworkPolicies (future)", "Traffic segmentation"),
        ("Identity", "Access Entries + RBAC", "Permission boundaries"),
        ("Storage", "PVC per namespace", "Data isolation")
    ]

    for layer, mechanism, purpose in layers:
        p = tf.add_paragraph()
        p.text = f"{layer}: {mechanism}"
        p.font.size = Pt(14)
        p.level = 1

        p = tf.add_paragraph()
        p.text = f"→ {purpose}"
        p.font.size = Pt(12)
        p.font.italic = True
        p.level = 2
        p.font.color.rgb = LIGHT_GRAY

    p = tf.add_paragraph()
    p.text = "\nNode Group Strategy:"
    p.font.size = Pt(18)
    p.font.bold = True

    strategy = [
        "System nodes: No taints → accept platform workloads",
        "Tenant nodes: customer-id=customerX:NoSchedule → dedicated compute",
        "Prevents 'noisy neighbor' issues",
        "Enables per-tenant node sizing/scaling"
    ]

    for item in strategy:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(13)
        p.level = 1

    # Slide 5: Access Control & Security
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Access Control & Security"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "IAM to Kubernetes Integration:"
    p.font.size = Pt(18)
    p.font.bold = True

    flow = """User/Team Role (AWS Access Account)
    ↓ AssumeRole
WorkloadConfig Role (Project Account)
    ↓ EKS Access Entry
Kubernetes Group (tenant:customer1)
    ↓ RoleBinding
ClusterRole (tenant-editor)
    ↓ Permissions
Namespace (tenant-customer1)"""

    p = tf.add_paragraph()
    p.text = flow
    p.font.size = Pt(13)
    p.font.name = 'Courier New'
    p.level = 1

    p = tf.add_paragraph()
    p.text = "\nSecurity Features:"
    p.font.size = Pt(18)
    p.font.bold = True

    features = [
        "✅ Pod Identity with Permissions Boundary",
        "✅ RBAC with least-privilege roles",
        "✅ Encrypted EBS volumes (GP3)",
        "✅ No shell access (pods/exec removed)",
        "✅ Read-only logs/events (audit trail)",
        "✅ ServiceAccount creation restricted"
    ]

    for item in features:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.level = 1
        p.font.color.rgb = GREEN

    # Slide 6: Add-ons & Platform Services
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Add-ons & Platform Services"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    addons = [
        ("Amazon EBS CSI Driver", "Persistent volumes for StatefulSets", "✅ Tested"),
        ("Amazon VPC CNI", "Pod networking", "✅ Active"),
        ("BRUPOP", "Automated Bottlerocket updates", "✅ Deployed"),
        ("cert-manager", "TLS certificate management", "✅ Running"),
        ("Pod Identity Agent", "IAM roles for pods", "✅ Configured")
    ]

    for addon, purpose, status in addons:
        p = tf.paragraphs[0] if addon == addons[0][0] else tf.add_paragraph()
        p.text = f"{addon}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = DARK_GRAY

        p = tf.add_paragraph()
        p.text = f"Purpose: {purpose}"
        p.font.size = Pt(12)
        p.level = 1

        p = tf.add_paragraph()
        p.text = f"Status: {status}"
        p.font.size = Pt(12)
        p.level = 1
        if "✅" in status:
            p.font.color.rgb = GREEN
        else:
            p.font.color.rgb = AWS_ORANGE

    p = tf.add_paragraph()
    p.text = "\nWhat This Enables:"
    p.font.size = Pt(18)
    p.font.bold = True

    enables = [
        "Databases with persistent storage",
        "Auto-healing infrastructure",
        "Secure AWS service access (S3, RDS, Secrets Manager)"
    ]

    for item in enables:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.level = 1

    # Slide 7: Demo Preview
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Demo - Flask App with Persistent Storage"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "What I'll Show:"
    p.font.size = Pt(20)
    p.font.bold = True

    demo_items = [
        ("1. Application Deployment", [
            "StatefulSet with 2 replicas",
            "Each pod gets dedicated 5GB EBS volume"
        ]),
        ("2. Persistence Test", [
            "Add data via REST API",
            "Delete pod (simulate failure)",
            "Pod recreates, data intact ✅"
        ]),
        ("3. Tenant Isolation", [
            "Pods scheduled only on customer1 nodes",
            "PVCs scoped to namespace",
            "RBAC prevents cross-namespace access"
        ])
    ]

    for section, items in demo_items:
        p = tf.add_paragraph()
        p.text = section
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = BLUE

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(14)
            p.level = 1

    p = tf.add_paragraph()
    p.text = "\n🎯 Key Takeaway: Workloads survive failures with zero data loss"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GREEN

    # Slide 8: What's Working Today
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "What's Working Today"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    working = [
        ("Compute & Scheduling", [
            "Multi-tenant node isolation",
            "StatefulSet deployments",
            "Pod scheduling with affinity/tolerations"
        ]),
        ("Storage", [
            "Dynamic EBS volume provisioning",
            "Volume persistence across pod restarts",
            "Encrypted volumes (GP3, 3000 IOPS baseline)"
        ]),
        ("Identity & Access", [
            "IAM-to-Kubernetes mapping",
            "Namespace-scoped permissions",
            "Pod Identity for AWS service access"
        ]),
        ("Operations", [
            "Automated OS updates (BRUPOP)",
            "GitOps-ready (manifests in repo)"
        ])
    ]

    for category, items in working:
        p = tf.paragraphs[0] if category == working[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = DARK_GRAY

        for item in items:
            p = tf.add_paragraph()
            p.text = f"✅ {item}"
            p.font.size = Pt(14)
            p.level = 1
            p.font.color.rgb = GREEN

    # Slide 9: Testing Completed
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Testing & Validation Completed"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    tests = [
        ("✅ Infrastructure Tests", [
            "Cluster deployment via CloudFormation",
            "Node group auto-scaling",
            "EBS CSI driver IAM permissions"
        ]),
        ("✅ Application Tests", [
            "Container deployment (Flask app)",
            "Persistent volume lifecycle",
            "Pod restart/failure scenarios",
            "Cross-AZ scheduling"
        ]),
        ("✅ Security Tests", [
            "RBAC permission boundaries",
            "Namespace isolation",
            "Pod Identity integration"
        ]),
        ("⏳ Pending Tests", [
            "Secrets Manager integration",
            "Network policies",
            "Resource quotas enforcement"
        ])
    ]

    for category, items in tests:
        p = tf.paragraphs[0] if category == tests[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(16)
        p.font.bold = True
        if "✅" in category:
            p.font.color.rgb = GREEN
        else:
            p.font.color.rgb = AWS_ORANGE

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(13)
            p.level = 1

    # Slide 10: Current Gaps
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Current Gaps & Limitations"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    gaps = [
        ("Observability", [
            "❌ No centralized logging",
            "❌ No metrics/monitoring dashboard",
            "❌ No alerting"
        ]),
        ("Environment Management", [
            "❌ Single cluster (no non-prod/prod split)",
            "❌ No CI/CD integration demonstrated"
        ]),
        ("IAM Permissions", [
            "⚠️  EKS Management role limited to eks:*",
            "Needs: ec2, iam, cloudformation, s3"
        ]),
        ("Testing Gaps", [
            "❌ Secrets Manager not tested",
            "❌ Network policies not implemented",
            "❌ Service mesh not evaluated"
        ])
    ]

    for category, items in gaps:
        p = tf.paragraphs[0] if category == gaps[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RED

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(13)
            p.level = 1
            if "⚠️" in item or "❌" in item:
                p.font.color.rgb = RED

    # Slide 11: Phase 1 Roadmap
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Future Roadmap - Phase 1 (Next 2-4 weeks)"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    phase1 = [
        ("1. Observability Stack", [
            "FluentBit/Fluentd → CloudWatch Logs",
            "Prometheus + Grafana → Metrics",
            "CloudWatch Container Insights"
        ]),
        ("2. Secrets Management", [
            "AWS Secrets Manager CSI Driver",
            "Mount secrets as volumes",
            "Auto-rotation support"
        ]),
        ("3. Environment Strategy", [
            "Decision: Separate clusters (non-prod, prod)",
            "Rationale: Blast radius control",
            "Same templates, different configs"
        ]),
        ("4. IAM Permission Expansion", [
            "eks:* (current) ✅",
            "ec2:*, iam:*, cloudformation:* 🔒",
            "BLOCKER: Needed for operations"
        ])
    ]

    for item_num, (category, details) in enumerate(phase1):
        p = tf.paragraphs[0] if item_num == 0 else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = BLUE

        for detail in details:
            p = tf.add_paragraph()
            p.text = detail
            p.font.size = Pt(12)
            p.level = 1
            if "BLOCKER" in detail:
                p.font.color.rgb = RED
                p.font.bold = True

    # Slide 12: Phase 2 Roadmap
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Future Roadmap - Phase 2 (1-3 months)"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    phase2 = [
        ("Network Isolation", [
            "Calico or Cilium for NetworkPolicies",
            "Deny-by-default traffic rules"
        ]),
        ("Resource Management", [
            "ResourceQuotas per namespace",
            "LimitRanges to prevent hogging",
            "Pod Disruption Budgets for HA"
        ]),
        ("Developer Experience", [
            "ArgoCD for GitOps deployments",
            "Kubernetes Dashboard (read-only)",
            "Self-service onboarding"
        ]),
        ("Compliance & Auditing", [
            "CloudTrail integration",
            "Kubernetes audit logs → SIEM",
            "Policy enforcement (Kyverno/OPA)"
        ])
    ]

    for category, items in phase2:
        p = tf.paragraphs[0] if category == phase2[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = BLUE

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(12)
            p.level = 1

    # Slide 13: Phase 3 Roadmap
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Future Roadmap - Phase 3 (3-6 months)"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    phase3 = [
        ("Service Mesh", [
            "Istio or Linkerd",
            "mTLS between services",
            "Advanced traffic management"
        ]),
        ("Cluster Autoscaler", [
            "Scale nodes based on pending pods",
            "Karpenter for better bin-packing"
        ]),
        ("Backup & DR", [
            "Velero for cluster backups",
            "Cross-region replication strategy"
        ]),
        ("Cost Optimization", [
            "Spot instances for non-critical workloads",
            "Kubecost for cost visibility",
            "Right-sizing recommendations"
        ])
    ]

    for category, items in phase3:
        p = tf.paragraphs[0] if category == phase3[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = BLUE

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(12)
            p.level = 1

    # Slide 14: Non-Prod vs Prod Strategy
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Non-Prod vs Prod Strategy"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "Recommended: Separate Clusters"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = GREEN

    comparison = [
        ("Availability", "Single AZ", "Multi-AZ"),
        ("Nodes", "Smaller instances", "Production-grade"),
        ("Backups", "Optional", "Automated"),
        ("Updates", "Test first", "After validation"),
        ("Access", "Broader (devs)", "Restricted (ops)"),
        ("Logging", "7-day retention", "90+ days"),
        ("Cost", "~30% of prod", "Full production")
    ]

    p = tf.add_paragraph()
    p.text = "\nAspect          Non-Prod               Prod"
    p.font.size = Pt(13)
    p.font.name = 'Courier New'
    p.font.bold = True

    for aspect, nonprod, prod in comparison:
        p = tf.add_paragraph()
        p.text = f"{aspect:15} {nonprod:22} {prod}"
        p.font.size = Pt(12)
        p.font.name = 'Courier New'
        p.level = 1

    p = tf.add_paragraph()
    p.text = "\nWhy Separate Clusters:"
    p.font.size = Pt(16)
    p.font.bold = True

    reasons = [
        "✅ Blast radius containment",
        "✅ Different SLAs/SLOs",
        "✅ Independent scaling",
        "✅ Compliance requirements"
    ]

    for reason in reasons:
        p = tf.add_paragraph()
        p.text = reason
        p.font.size = Pt(13)
        p.level = 1
        p.font.color.rgb = GREEN

    # Slide 15: Key Decisions Needed
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Key Decisions Needed"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    decisions = [
        ("Decision 1: Environment Strategy",
         "Separate non-prod/prod clusters (recommended)",
         "Needed before onboarding prod workloads",
         RED),
        ("Decision 2: Logging Destination",
         "CloudWatch Logs / OpenSearch / Third-party",
         "Critical for troubleshooting",
         RED),
        ("Decision 3: IAM Permission Expansion",
         "Approve expanded EKS Management role",
         "BLOCKER: Current eks:* too restrictive",
         RED),
        ("Decision 4: Network Isolation",
         "Implement NetworkPolicies now or defer",
         "Before sensitive workloads",
         AWS_ORANGE)
    ]

    for i, (decision, option, timeline, color) in enumerate(decisions):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = decision
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = color

        p = tf.add_paragraph()
        p.text = f"☐ {option}"
        p.font.size = Pt(12)
        p.level = 1

        p = tf.add_paragraph()
        p.text = f"Timeline: {timeline}"
        p.font.size = Pt(11)
        p.font.italic = True
        p.level = 2
        p.font.color.rgb = LIGHT_GRAY

    # Slide 16: Recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Recommendations"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    recommendations = [
        ("Immediate (This Sprint)", [
            "✅ Document architecture (done)",
            "Expand EKS Management IAM permissions",
            "Test Secrets Manager CSI driver",
            "Set up CloudWatch Container Insights"
        ]),
        ("Next Sprint", [
            "Deploy FluentBit for logging",
            "Create non-prod cluster",
            "Implement ResourceQuotas",
            "Add Prometheus/Grafana"
        ]),
        ("Within 2 Months", [
            "NetworkPolicies implementation",
            "CI/CD pipeline integration",
            "Onboard first production tenant",
            "Disaster recovery testing"
        ])
    ]

    for category, items in recommendations:
        p = tf.paragraphs[0] if category == recommendations[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = BLUE

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(13)
            p.level = 1
            if "✅" in item:
                p.font.color.rgb = GREEN

    # Slide 17: Success Metrics
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Success Metrics"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    metrics = [
        ("Technical Metrics", [
            "Cluster uptime: 99.9% target",
            "Pod startup time: <30 seconds",
            "PVC provisioning: <2 minutes",
            "Zero cross-tenant security incidents"
        ]),
        ("Business Metrics", [
            "Cost per tenant: $200-270/month",
            "Time to onboard tenant: <1 day",
            "Developer satisfaction score",
            "Reduction in infrastructure tickets"
        ]),
        ("Operational Metrics", [
            "Mean time to recovery (MTTR)",
            "Incident count per month",
            "Automated vs manual operations ratio"
        ])
    ]

    for category, items in metrics:
        p = tf.paragraphs[0] if category == metrics[0][0] else tf.add_paragraph()
        p.text = category
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = DARK_GRAY

        for item in items:
            p = tf.add_paragraph()
            p.text = item
            p.font.size = Pt(13)
            p.level = 1

    # Slide 18: Risk & Mitigation
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Risk & Mitigation"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    risks = [
        ("Single cluster outage", "High", "Multi-AZ + DR cluster plan"),
        ("Tenant escapes namespace", "Critical", "Security audits + NetworkPolicies"),
        ("EBS volume failure", "Medium", "Backup strategy + replicas"),
        ("IAM permission gaps", "Medium", "Expand EKS Management role"),
        ("Cost overruns", "Medium", "ResourceQuotas + monitoring"),
        ("Lack of observability", "High", "FluentBit + CloudWatch (priority)")
    ]

    p = tf.paragraphs[0]
    p.text = "Risk                        Impact    Mitigation"
    p.font.size = Pt(13)
    p.font.name = 'Courier New'
    p.font.bold = True

    for risk, impact, mitigation in risks:
        p = tf.add_paragraph()
        p.text = f"{risk:28} {impact:8} {mitigation}"
        p.font.size = Pt(11)
        p.font.name = 'Courier New'
        p.level = 1

        if impact == "Critical":
            p.font.color.rgb = RED
        elif impact == "High":
            p.font.color.rgb = AWS_ORANGE

    # Slide 19: Q&A
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    left = Inches(1)
    top = Inches(2.5)
    width = Inches(8)
    height = Inches(2)

    qa_box = slide.shapes.add_textbox(left, top, width, height)
    qa_frame = qa_box.text_frame
    qa_frame.text = "Questions & Discussion"
    qa_para = qa_frame.paragraphs[0]
    qa_para.font.size = Pt(48)
    qa_para.font.bold = True
    qa_para.font.color.rgb = DARK_GRAY
    qa_para.alignment = PP_ALIGN.CENTER

    topics_box = slide.shapes.add_textbox(Inches(2), Inches(4.5), Inches(6), Inches(2))
    topics_frame = topics_box.text_frame
    topics_text = """Open Questions:
• Preferred logging solution?
• Timeline for prod workload onboarding?
• Budget for observability tools?

Demo Ready:
• Live cluster walkthrough
• Persistence demonstration"""

    topics_frame.text = topics_text
    for para in topics_frame.paragraphs:
        para.font.size = Pt(16)
        para.alignment = PP_ALIGN.CENTER
        para.font.color.rgb = LIGHT_GRAY

    # Slide 20: Backup - Cost Breakdown
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "BACKUP: Cost Breakdown"

    content = slide.placeholders[1]
    tf = content.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.text = "Current Monthly Cost Estimate (Single Tenant):"
    p.font.size = Pt(18)
    p.font.bold = True

    costs = [
        ("EKS Control Plane", "$73"),
        ("System nodes (2x m5.large)", "~$140"),
        ("Customer nodes (1x m5.xlarge)", "~$140"),
        ("EBS volumes (10GB)", "~$1"),
        ("Data transfer", "~$50"),
        ("", ""),
        ("Total", "~$404/month")
    ]

    for item, cost in costs:
        p = tf.add_paragraph()
        if item == "Total":
            p.text = f"\n{item:35} {cost}"
            p.font.bold = True
            p.font.size = Pt(16)
            p.font.color.rgb = GREEN
        elif item == "":
            p.text = "─" * 50
            p.font.size = Pt(14)
        else:
            p.text = f"{item:35} {cost}"
            p.font.size = Pt(14)
            p.font.name = 'Courier New'
        p.level = 1

    p = tf.add_paragraph()
    p.text = "\nWith 3 tenants: ~$600-800/month"
    p.font.size = Pt(16)
    p.font.bold = True

    p = tf.add_paragraph()
    p.text = "Cost per tenant: $200-270/month"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = GREEN

    # Save presentation
    output_path = "/Users/manavkhanna/Documents/Projects/Kubernetes/EKS-Multi-Tenant-Cluster-Presentation.pptx"
    prs.save(output_path)
    return output_path

if __name__ == "__main__":
    try:
        output = create_presentation()
        print(f"✅ PowerPoint created successfully: {output}")
    except Exception as e:
        print(f"❌ Error creating PowerPoint: {e}")
        import traceback
        traceback.print_exc()
